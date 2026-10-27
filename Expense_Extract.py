# Expense Report Extractor
# Written by Weston Matthews (2020)

# Takes a date range, and selects a QuickBooks company file
# Extracts a Profit&Loss report from the QuickBooks company file, and outputs it as an Excel spreadsheet
# Then, edits the P&L report spreadsheet, removing Sales, COGS, and other income
# Outputs the resulting expense report

# Utilizes the EasyGui package for file selection handling
import easygui as gui

# Utilizes the datetime package for date object handling
from datetime import datetime

# Utilizes the PyWin32 COM wrapper to talk to the QBFC COM listener
import win32com.client as win32

import __main__

import os

import openpyxl as xl

# AppID and AppName, for communicating with QB
appID = "EXPENSE_EXTRACT_v1"
appName = "Expense Report Extractor v1.0"

def check_date(date):
    if type(date) != str:
        return -1
    else:
        if date.count('/') == 2:
            slash1 = date.find('/')
            slash2 = date.rfind('/')
            month = date[:slash1]
            day = date[slash1+1:slash2]
            year = date[slash2+1:]
            try:
                month = int(month)
                day = int(day)
            except:
                return -1
            if len(year) == 2:
                year = "20"+year
                try:
                    year = int(year)
                except:
                    return -1
            elif len(year) == 4:
                try:
                    year = int(year)
                except:
                    return -1
            else:
                return -1
            return_date = datetime(year, month, day)
            return return_date
        else:
            return -1

def calc_ws_col(rowTypes):
    # Returns the furthest column of the headers, to the left and the right
    max_col = 0
    min_col = 0
    col = 0
    lastType = rowTypes[0]
    for row in rowTypes[1:]:
        if row == 'h':
            if lastType == 'h':
                col += 1
            lastType = 'h'
        elif row == 'r':
            if lastType == 'h':
                col += 1
            elif lastType == 's':
                col += 1
            lastType = 'r'
        elif row == 's':
            col -= 1
            lastType = 's'
        elif row == 't':
            col -= 1
            lastType = 't'
        if col > max_col:
            max_col = col
        if col < min_col:
            min_col = col
        #print(str(col))
    return (min_col, max_col)

def extract_expense(rowData):
    #Find the beginning of the expense report
    start = 0
    end = 0
    for row in rowData:
        if type(row) == str:
            if row == "Expense":
                start = rowData.index(row)
            elif row == "Total Expense":
                end = rowData.index(row)
        elif type(row) == list:
            if row[0][1] == "Expense":
                start = rowData.index(row)
            elif row[0][1] == "Total Expense":
                end = rowData.index(row)
    return (rowData[start:end+1], start, end)

def main():
    # Memory initialization
    return_log = ""

    # Input the date range
    msg = "Enter the date range to extract.\n(Must be in the form MM/DD/YY or MM/DD/YYYY)"
    title = "Expense Report Extractor v1.0"
    fieldNames = ["Start Date", "End Date"]
    fieldValues = []
    fieldValues = gui.multenterbox(msg, title, fieldNames)
    if fieldValues == None:
        return 0
    else:
        start_date = fieldValues[0]
        end_date = fieldValues[1]
        start_date = check_date(start_date)
        end_date = check_date(end_date)
        if start_date == -1:
            return_log = ""
            return_log += "ERROR!\n"
            return_log += "Invalid start date entered!\n"
            return_log += "You entered: "+str(fieldValues[0])+"\n"
            msg = "Here's what happened:"
            gui.textbox(msg, title, text=return_log)
            return 0
        if end_date == -1:
            return_log = ""
            return_log += "ERROR!\n"
            return_log += "Invalid end date entered!\n"
            return_log += "You entered: "+str(fieldValues[1])+"\n"
            msg = "Here's what happened:"
            gui.textbox(msg, title, text=return_log)
            return 0
        print(start_date)
        print(end_date)
    
    # Select the QB company file
    msg = "Please select the QB company file"
    title = "Expense Report Extractor v1.0"
    gui.msgbox(msg, title)
    file = gui.fileopenbox(msg, title)
    
    company = file[file.rfind('\\'):]

    # Open a QB Session
    sessionManager = win32.DispatchEx("QBFC13.QBSessionManager")
    print("QuickBooks opened.\n")
    
    # Open a connection to QB
    sessionManager.OpenConnection(appID, appName)
    print("Connection opened.\n")
    
    # Begin Session
    try:
        sessionManager.BeginSession(file, 0)
        print("Session started.\n")
    except:
        gui.exceptionbox(title=title)
        return 0
    
    # Construct request message set
    requestMsgSet = sessionManager.CreateMsgSetRequest("US", 7, 0)
    print("Request Message Set created.\n")
    
    # Requesting a GeneralSummaryReportQuery
    summaryReportRq = requestMsgSet.AppendGeneralSummaryReportQueryRq()
    print("Requesting General Summary Report Query.\n")
    
    # Selecting the Profit and Loss Summary report
    summaryReportRq.GeneralSummaryReportType.SetAsString("ProfitAndLossStandard")
    print("Selected Profit and Loss Standard report type.\n")
    
    # Inputing the info for the P&L report
    summaryReportRq.ORReportPeriod.ReportPeriod.FromReportDate.SetValue(start_date)
    summaryReportRq.ORReportPeriod.ReportPeriod.ToReportDate.SetValue(end_date)
    summaryReportRq.SummarizeColumnsBy.SetAsString("Month")
    
    # Send Request message set
    responseMsgSet = sessionManager.DoRequests(requestMsgSet)
    print("Request Set sent.\n")
    
    # Disconnect from Quickbooks
    sessionManager.EndSession()           # Close the company file
    print("Session ended. Company file closed.\n")
    sessionManager.CloseConnection()      # Close the connection
    print("Connection closed.\n")
    
    # Response handling
    # Grab the response message set
    QBXML = responseMsgSet
    QBXMLMsgsRq = QBXML.ResponseList
    QueryRs = QBXMLMsgsRq.GetAt(0)
    print("Printing GeneralSummaryReportQuery response information:\n")
    statusCode = QueryRs.StatusCode
    statusSeverity = QueryRs.StatusSeverity
    statusMessage = QueryRs.StatusMessage
    print("Status Code: "+str(statusCode))
    print("Status Severity: "+str(statusSeverity))
    print("Status Message: "+str(statusMessage))
    print("\n")
    
    if statusSeverity == "Info":
        now = datetime.now()
    
        # Begin printing info from the report query
        reportTitle = QueryRs.Detail.ReportTitle.GetValue()
        print("Report Title: "+str(reportTitle))
        reportSubTitle = QueryRs.Detail.ReportSubtitle.GetValue()
        print("Report Subtitle: "+str(reportSubTitle))
        reportBasis = QueryRs.Detail.ReportBasis.GetAsString()
        print("Report Basis: "+str(reportBasis))
        reportNumRows = QueryRs.Detail.NumRows.GetValue()
        print("NumRows: "+str(reportNumRows))
        reportNumCols = QueryRs.Detail.NumColumns.GetValue()
        print("NumColumns: "+str(reportNumCols))
        reportNumTitleRows = QueryRs.Detail.NumColTitleRows.GetValue()
        print("NumColTitleRows: "+str(reportNumTitleRows))
        
        colDescList = QueryRs.Detail.ColDescList
        titleCols = []
        for x in range(0, len(colDescList)):
            col = colDescList.GetAt(x)
            for y in range(0, len(col.ColTitleList)):
                title = col.ColTitleList.GetAt(y)
                if title.value != None:
                    #print(title.value.GetValue())
                    titleCols.append(title.value.GetValue())
                else:
                    #print("None")
                    titleCols.append('')
        print(titleCols)
        
        rowVals = []
        rowTypes = []
        reportDataList = QueryRs.Detail.ReportData.ORReportDataList
        print(len(reportDataList))
        for x in range(0, len(reportDataList)):
            reportData = reportDataList.GetAt(x)
            if reportData.DataRow != None:
                if reportData.DataRow.RowData == None:
                    cellData = []
                    cellData.append((reportData.DataRow.ColDataList.GetAt(0).colID.GetValue(), reportData.DataRow.ColDataList.GetAt(0).value.GetValue()))
                    for z in range(2, reportNumCols+1):
                        cellData.append((z, '0.00'))
                    rowVals.append(cellData)
                    rowTypes.append('r')
                elif reportData.DataRow.ColDataList != None:
                    colDataList = reportData.DataRow.ColDataList
                    cellData = []
                    for y in range(0, len(colDataList)):
                        cellData.append((colDataList.GetAt(y).colID.GetValue(), colDataList.GetAt(y).value.GetValue()))
                    rowVals.append(cellData)
                    rowTypes.append('r')
            elif reportData.TextRow != None:
                cellData = reportData.TextRow.value.GetValue()
                rowVals.append(cellData)
                rowTypes.append('h')
            elif reportData.SubtotalRow != None:
                if reportData.SubtotalRow.RowData == None:
                    cellData = []
                    cellData.append((reportData.SubtotalRow.ColDataList.GetAt(0).colID.GetValue(), reportData.SubtotalRow.ColDataList.GetAt(0).value.GetValue()))
                    for z in range(2, reportNumCols+1):
                        cellData.append((z, '0.00'))
                    rowVals.append(cellData)
                    rowTypes.append('s')
                elif reportData.SubtotalRow.ColDataList != None:
                    colDataList = reportData.SubtotalRow.ColDataList
                    cellData = []
                    for y in range(0, len(colDataList)):
                        cellData.append((colDataList.GetAt(y).colID.GetValue(), colDataList.GetAt(y).value.GetValue()))
                    rowVals.append(cellData)
                    rowTypes.append('s')
            elif reportData.TotalRow != None:
                if reportData.TotalRow.RowData == None:
                    cellData = []
                    cellData.append((reportData.TotalRow.ColDataList.GetAt(0).colID.GetValue(), reportData.TotalRow.ColDataList.GetAt(0).value.GetValue()))
                    for z in range(2, reportNumCols+1):
                        cellData.append((z, '0.00'))
                    rowVals.append(cellData)
                    rowTypes.append('t')
                elif reportData.TotalRow.ColDataList != None:
                    colDataList = reportData.TotalRow.ColDataList
                    cellData = []
                    for y in range(0, len(colDataList)):
                        cellData.append((colDataList.GetAt(y).colID.GetValue(), colDataList.GetAt(y).value.GetValue()))
                    rowVals.append(cellData)
                    rowTypes.append('t')
        #print(rowTypes)
        #for row in rowVals:
            #print(row)
        
        ## Begin printing the data into an excel spreadsheet
        # Create a new excel spreadsheet workbook
        wb = xl.Workbook()
        # Select the primary page and name it "Report"
        ws = wb.active
        ws.title = "Report"
        
        # Initialization of memory
        i_row = 2
        i_col = 2
        cell_data = ""
        extract_result = extract_expense(rowVals)
        rowVals = extract_result[0]
        types_start = extract_result[1]
        types_end = extract_result[2]
        rowTypes = rowTypes[types_start:types_end+1]
        print(rowTypes)
        for row in rowVals:
            print(row)
        head_col_range = calc_ws_col(rowTypes)
        max_head_col = head_col_range[1]
        min_head_col = head_col_range[0]
        print("Max col: "+str(max_head_col))
        print("Min col: "+str(min_head_col))
        if min_head_col <= 0:
            i_col += -1*min_head_col
        
        # Constructing the cell styles
        cellFont = xl.styles.Font(name='Arial', size=8)
        cellFont_B = xl.styles.Font(name='Arial', size=8, bold=True)
        side_thin = xl.styles.Side(border_style="thin", color='000000')
        side_thick = xl.styles.Side(border_style="thick", color='000000')
        cellAlign = xl.styles.Alignment(horizontal="center", vertical="center")
        acctFormat = '#,##0.00;\\-#,##0.00'
        
        # Printing header and footer
        header_text = ""
        if now.hour > 12:
            header_text += str(now.hour-12)
        else:
            header_text += str(now.hour)
        header_text += ":"+str(now.minute)
        if now.hour > 12:
            header_text += " PM"
        else:
            header_text += " AM"
        header_text += "\n"+str(now.month)+"/"+str(now.day)+"/"+str(now.year)+"\n"
        header_text += reportBasis+" Basis"
        ws.oddHeader.left.text = header_text
        ws.oddHeader.left.size = 8
        ws.oddHeader.left.font = "Arial"
        ws.oddHeader.left.bold = True
        header_text = "Expense Report"
        ws.oddHeader.center.text = header_text
        ws.oddHeader.center.size = 14
        ws.oddHeader.center.bold = True
        ws.oddHeader.center.font = "Arial"
        
        # Begin printing row data
        if rowTypes[0] == 'h':
            rowDat = rowVals[0]
            lastType = 'h'
            cell = ws.cell(column=i_col, row=i_row, value=rowDat)
            cell.font = cellFont_B
        elif rowTypes[0] == 'r':
            rowDat = rowVals[0][1]
            lastType = 'r'
            cell = ws.cell(column=i_col, row=i_row, value=rowDat)
            cell.font = cellFont_B
        for x in range(1, len(rowTypes)):
            rowType = rowTypes[x]
            rowDat = rowVals[x]
            if rowType == 'h':
                if lastType == 'h':
                    i_row += 1
                    i_col += 1
                    cell = ws.cell(column=i_col, row=i_row, value=rowDat)
                    cell.font = cellFont_B
                else:
                    i_row += 1
                    cell = ws.cell(column=i_col, row=i_row, value=rowDat)
                    cell.font = cellFont_B
                lastType = 'h'
            elif rowType == 'r':
                if lastType == 'h':
                    i_row += 1
                    i_col += 1
                    cell = ws.cell(column=i_col, row=i_row, value=rowDat[0][1])
                    cell.font = cellFont_B
                else:
                    i_row += 1
                    cell = ws.cell(column=i_col, row=i_row, value=rowDat[0][1])
                    cell.font = cellFont_B
                lastType = 'r'
            elif rowType == 's':
                i_row += 1
                i_col -= 1
                cell = ws.cell(column=i_col, row=i_row, value=rowDat[0][1])
                cell.font = cellFont_B
                lastType = 's'
            elif rowType == 't':
                i_row += 1
                i_col -= 1
                cell = ws.cell(column=i_col, row=i_row, value=rowDat[0][1])
                cell.font = cellFont_B
                lastType = 't'
        
        # Print the Header row
        i_row = 1
        i_col = 2+max_head_col
        if titleCols[0] == '':
            titleCols = titleCols[1:]
        for title in titleCols:
            cell = ws.cell(column=i_col, row=i_row, value=title)
            cell.font = cellFont_B
            cell.alignment = cellAlign
            cell.border = xl.styles.Border(bottom=side_thick)
            i_col += 2
        
        # Print cell data/formulas into cells
        i_row = 3
        i_col = 2+max_head_col
        cur_subtotals = []
        tot_subtotals = []
        tot_cells = []
        subtot_start = ""
        subtot_end = ""
        lastType = ""
        final_row = len(rowTypes)-1
        for x in range(1, len(rowTypes)):
            rowType = rowTypes[x]
            rowDat = rowVals[x]
            i_col = 2+max_head_col
            for y in range(1, len(titleCols)+1):
                if type(rowDat) != str:
                    cellDat = rowDat[y]
                    #print(cellDat)
                if rowType == 'r':
                    cell = ws.cell(column=i_col, row=i_row, value=float(cellDat[1]))
                    cell.font = cellFont
                    cell.number_format = acctFormat
                    print(str(rowType)+" Cell ("+str(i_row)+", "+str(i_col)+") <- "+str(cellDat[1]))
                elif rowType == 's':
                    if x == final_row:
                        # Calculate the total of all sub-header amounts
                        # "Expense" is at (2,2), so, we start looking at (3,3) and move down row by row
                        j_row = 3
                        j_col = 3
                        tot_cells = []
                        while j_row < x:
                            if ws.cell(column=3, row=j_row).value != "" and ws.cell(column=3, row=j_row).value != None:
                                if ws.cell(column=i_col, row=j_row).value != "" and ws.cell(column=i_col, row=j_row).value != None:
                                    tot_cells.append(ws.cell(column=i_col, row=j_row).column_letter+str(j_row))
                            j_row += 1
                        tot_string = ""
                        for cell in tot_cells:
                            tot_string += cell+","
                        tot_string = tot_string[:len(tot_string)-1]
                        if tot_string != "":
                            cell = ws.cell(column=i_col, row=i_row, value="=SUM("+tot_string+")")
                            cell.font = cellFont
                            cell.number_format = acctFormat
                            cell.border = xl.styles.Border(top=side_thin, bottom=side_thin)
                            print(str(rowType)+" Cell ("+str(i_row)+", "+str(i_col)+") <- "+"=SUM("+tot_string+")")
                    elif lastType == 'r':
                        # Sum each cell above the subtotal cell
                        j_row = i_row-1
                        subtot_end = ws.cell(column=i_col, row=j_row).column_letter+str(j_row)
                        while j_row > 0:
                            j_row -= 1
                            if ws.cell(column=i_col, row=j_row).value == '' or ws.cell(column=i_col, row=j_row).value == None:
                                subtot_start = ws.cell(column=i_col, row=j_row+1).column_letter+str(j_row+1)
                                break
                        cell = ws.cell(column=i_col, row=i_row, value="=ROUND(SUM("+subtot_start+":"+subtot_end+"),5)")
                        cell.font = cellFont
                        cell.number_format = acctFormat
                        cell.border = xl.styles.Border(top=side_thin)
                        print(str(rowType)+" Cell ("+str(i_row)+", "+str(y)+") <- "+"=ROUND(SUM("+subtot_start+":"+subtot_end+"),5)")
                        cur_subtotals.append(subtot_end[0]+str(i_row))
                    elif lastType == 'h':
                        # For some reason, a subtotal came after a header
                        cell = ws.cell(column=i_col, row=i_row, value=0.00)
                        cell.font = cellFont
                        cell.number_format = acctFormat
                        print(str(rowType)+" Cell ("+str(i_row)+", "+str(y)+") <- "+str(0.00))
                    elif lastType == 's':
                        # Subtotal of a subtotal
                        if len(cur_subtotals) == 1:
                            cell = ws.cell(column=i_col, row=i_row, value="="+cur_subtotals[0])
                            cell.font = cellFont
                            cell.border = xl.styles.Border(top=side_thin)
                            cell.number_format = acctFormat
                            print(str(rowType)+" Cell ("+str(i_row)+", "+str(y)+") <- "+"="+cur_subtotals[0])
                        elif len(cur_subtotals) > 1:
                            cell = ws.cell(column=i_col, row=i_row, value="=ROUND(SUM("+cur_subtotals[0]+":"+cur_subtotals[-1]+"),5)")
                            cell.font = cellFont
                            cell.border = xl.styles.Border(top=side_thin)
                            cell.number_format = acctFormat
                            print(str(rowType)+" Cell ("+str(i_row)+", "+str(y)+") <- "+"=ROUND(SUM("+cur_subtotals[0]+":"+cur_subtotals[-1]+"),5)")
                        cur_subtotals.append(ws.cell(column=i_col, row=i_row).column_letter+str(i_row))
                elif rowType == 'h':
                    pass
                elif rowType == 't':
                    pass
                i_col += 2
            i_row += 1
            lastType = rowType
        
        # Cleaning up column widths
        for x in range(1, 2+max_head_col-1):
            column_letter = xl.utils.cell.get_column_letter(x)
            ws.column_dimensions[column_letter].width = 2.29
        column_letter = xl.utils.cell.get_column_letter(2+max_head_col-1)
        ws.column_dimensions[column_letter].width = 31.43
        i = 1
        maxcols = 2+max_head_col+(2*len(titleCols))
        for x in range(2+max_head_col, maxcols):
            column_letter = xl.utils.cell.get_column_letter(x)
            if i%2 != 0:
                ws.column_dimensions[column_letter].width = 11.00
            else:
                ws.column_dimensions[column_letter].width = 1.57
            i += 1
        
        cwd = os.getcwd()
        s_date = str(start_date.month)+"-"+str(start_date.day)+"-"+str(start_date.year)
        e_date = str(end_date.month)+"-"+str(end_date.day)+"-"+str(end_date.year)
        wb.save(cwd+"\\"+company+" Exp Report - "+s_date+" to "+e_date+".xlsx")
        
        return 0
        
    elif statusSeverity == "Warning":
        pass
    elif statusSeverity == "Error":
        pass
    
    print("All Done!\n")

if __name__ == "__main__":
    main()